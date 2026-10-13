"""
DOCX builder for review results (SDD §F4).

Spec:
  - Times New Roman 12 pt, justified, bordered table, auto section numbering.
  - No colour, no emoji.
  - Disclaimer at top AND bottom.
  - parse_ok=False → raw text with same styling.
  - Number of Rujukan in docx == number in parsed result.
  - File built on-the-fly (BytesIO), never written to disk.
  - Filename: Telaah-Review-YYYYMMDD-HHmm.docx
"""
from __future__ import annotations

import io
from datetime import UTC, datetime

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from ..models.schemas import DebateType, ParsedReport, ReviewerResult

# ── Styling constants ──────────────────────────────────────────────────────────

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(12)
DISCLAIMER_TEXT = (
    "DISCLAIMER: Dokumen ini merupakan asistensi akademik yang dihasilkan oleh "
    "sistem Telaah. Tidak menggantikan penilaian dosen pembimbing/penguji resmi."
)

REVIEWER_LABELS = {
    "metodologi": "Reviewer Metodologi",
    "orisinalitas": "Reviewer Orisinalitas",
    "revisi": "Reviewer Revisi",
}

DEBATE_TYPE_LABELS = {
    DebateType.kontradiksi: "KONTRADIKSI",
    DebateType.penguatan_silang: "PENGUATAN SILANG",
    DebateType.isu_unik: "ISU UNIK",
}


# ── Public API ─────────────────────────────────────────────────────────────────

def build_docx(report: ParsedReport, raw_output: str | None = None) -> bytes:
    """
    Build a .docx document from a ParsedReport.
    If parse_ok=False, falls back to a plain-text version of raw_output.
    Returns raw bytes of the .docx file.
    """
    doc = Document()
    _set_default_style(doc)

    _add_disclaimer(doc, section_num=None)
    _add_blank(doc)

    if not report.parse_ok:
        _build_raw_fallback(doc, raw_output)
    else:
        _build_structured(doc, report)

    _add_blank(doc)
    _add_disclaimer(doc, section_num=None)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_filename() -> str:
    """Return timestamped filename: Telaah-Review-YYYYMMDD-HHmm.docx"""
    ts = datetime.now(UTC).strftime("%Y%m%d-%H%M")
    return f"Telaah-Review-{ts}.docx"


# ── Internal helpers ───────────────────────────────────────────────────────────

def _set_default_style(doc: DocumentType) -> None:
    """Set Normal style to Times New Roman 12, justified."""
    style = doc.styles["Normal"]
    font = style.font
    font.name = FONT_NAME
    font.size = FONT_SIZE
    pf = style.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _add_blank(doc: DocumentType) -> None:
    p = doc.add_paragraph()
    _style_paragraph(p)


def _style_paragraph(p: Paragraph, bold: bool = False, size: int | None = None) -> None:
    """Apply font to all runs in a paragraph."""
    for run in p.runs:
        run.font.name = FONT_NAME
        run.font.size = Pt(size) if size else FONT_SIZE
        if bold:
            run.bold = True
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _add_heading(doc: DocumentType, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.name = FONT_NAME
    run.font.size = FONT_SIZE
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _add_body(doc: DocumentType, text: str) -> None:
    p = doc.add_paragraph(text)
    _style_paragraph(p)


def _add_disclaimer(doc: DocumentType, section_num: int | None) -> None:
    p = doc.add_paragraph()
    run = p.add_run(DISCLAIMER_TEXT)
    run.bold = True
    run.italic = True
    run.font.name = FONT_NAME
    run.font.size = FONT_SIZE
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _build_raw_fallback(doc: DocumentType, raw_output: str | None) -> None:
    _add_heading(doc, "HASIL REVIEW (Format Tidak Terstruktur)")
    _add_blank(doc)
    text = raw_output or "(Tidak ada output tersedia.)"
    # Split by lines to preserve paragraph breaks
    for line in text.splitlines():
        p = doc.add_paragraph(line)
        _style_paragraph(p)


def _build_structured(doc: DocumentType, report: ParsedReport) -> None:
    sec = 1  # section counter

    # ── Titik Perdebatan ───────────────────────────────────────────────────────
    if report.debate_points:
        _add_heading(doc, f"{sec}. TITIK PERDEBATAN")
        sec += 1
        for i, dp in enumerate(report.debate_points, 1):
            type_label = DEBATE_TYPE_LABELS.get(dp.type, dp.type.value)
            _add_heading(doc, f"  {i}. [{type_label}] {dp.aspect}")
            if dp.description:
                _add_body(doc, dp.description)
            if dp.reviewers:
                for line in dp.reviewers:
                    _add_body(doc, f"    {line}")
        _add_blank(doc)

    # ── Reviewer sections ──────────────────────────────────────────────────────
    for reviewer in report.reviewers:
        label = REVIEWER_LABELS.get(reviewer.key.value, reviewer.key.value.upper())
        _add_heading(doc, f"{sec}. {label.upper()}")
        sec += 1

        if reviewer.status.value in ("failed", "missing"):
            status_text = (
                "Data reviewer ini tidak tersedia."
                if reviewer.status.value == "missing"
                else "Terjadi kesalahan saat memproses reviewer ini."
            )
            _add_body(doc, status_text)
            _add_blank(doc)
            continue

        # Verdict
        if reviewer.verdict:
            _add_heading(doc, f"  Verdict: {reviewer.verdict}")

        # Ringkasan
        if reviewer.summary:
            _add_heading(doc, "  Ringkasan:")
            _add_body(doc, reviewer.summary)

        # Temuan — bordered table
        if reviewer.findings:
            _add_heading(doc, "  Temuan:")
            _add_blank(doc)
            _build_findings_table(doc, reviewer)

        # Rekomendasi Prioritas
        if reviewer.priority_recommendations:
            _add_heading(doc, "  Rekomendasi Prioritas:")
            for j, rec in enumerate(reviewer.priority_recommendations, 1):
                _add_body(doc, f"  {j}. {rec}")

        _add_blank(doc)


def _build_findings_table(doc: DocumentType, reviewer: ReviewerResult) -> None:
    """
    Build a bordered 3-column table: No | Aspek/Deskripsi | Rujukan.
    Rujukan count must equal len(findings) — even null ones get a cell.
    """
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"  # bordered

    # Header row
    hdr = table.rows[0].cells
    for cell, text in zip(hdr, ["No.", "Aspek & Deskripsi", "Rujukan"], strict=False):
        p = cell.paragraphs[0]
        run = p.add_run(text)
        run.bold = True
        run.font.name = FONT_NAME
        run.font.size = FONT_SIZE
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Data rows
    for finding in reviewer.findings:
        row = table.add_row().cells

        # Col 0: number
        p0 = row[0].paragraphs[0]
        r0 = p0.add_run(str(finding.no))
        r0.font.name = FONT_NAME
        r0.font.size = FONT_SIZE

        # Col 1: [ASPEK: ...] + description
        p1 = row[1].paragraphs[0]
        r1a = p1.add_run(f"[ASPEK: {finding.aspect}]\n")
        r1a.bold = True
        r1a.font.name = FONT_NAME
        r1a.font.size = FONT_SIZE
        r1b = p1.add_run(finding.description)
        r1b.font.name = FONT_NAME
        r1b.font.size = FONT_SIZE
        p1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        # Col 2: Rujukan verbatim (or "Tanpa rujukan")
        p2 = row[2].paragraphs[0]
        ref_text = finding.reference if finding.reference is not None else "Tanpa rujukan"
        r2 = p2.add_run(ref_text)
        r2.font.name = FONT_NAME
        r2.font.size = FONT_SIZE
        if finding.reference is None:
            r2.italic = True
        p2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
