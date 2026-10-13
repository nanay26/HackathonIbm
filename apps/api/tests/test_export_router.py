"""
Tests for POST /api/v1/export/docx (stateless export, SDD F5 / spec Bagian 9).

Covers:
  - Valid ParsedReport → 200, valid .docx, correct Content-Type
  - raw_output only → 200, valid .docx containing raw text
  - Rujukan count in docx == Rujukan in parsed (F4-4 parity)
  - Empty body (neither parsed nor raw_output) → 400
"""
from __future__ import annotations

import io

from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.models.schemas import (
    DebatePoint,
    DebateType,
    Finding,
    ParsedReport,
    ReviewerKey,
    ReviewerResult,
    ReviewerStatus,
)

client = TestClient(app, raise_server_exceptions=False)

# ── Fixtures ──────────────────────────────────────────────────────────────────

FINDINGS = [
    Finding(
        no=1,
        aspect="Desain Penelitian",
        description="Skema validasi tidak dijelaskan.",
        reference='"validasi 10-fold CV" (BAB 3, hlm. 24)',
    ),
    Finding(
        no=2,
        aspect="Reproduksibilitas",
        description="Hyperparameter tidak dicantumkan.",
        reference=None,
    ),
]

REVIEWER = ReviewerResult(
    key=ReviewerKey.metodologi,
    status=ReviewerStatus.ok,
    verdict="REVISI MAYOR",
    summary="Metodologi perlu perbaikan.",
    findings=FINDINGS,
    priority_recommendations=["Perbaiki skema validasi."],
)

DEBATE = DebatePoint(
    type=DebateType.penguatan_silang,
    aspect="Validitas",
    description="Dua reviewer menyoroti isu ini.",
    reviewers=["- Reviewer Metodologi: potensi data leakage."],
    references=[],
)

PARSED_OK = ParsedReport(
    parse_ok=True,
    reviewers=[REVIEWER],
    debate_points=[DEBATE],
)

PARSED_FAIL = ParsedReport(parse_ok=False, reviewers=[], debate_points=[])

RAW = "Output mentah dari Flow 5 sebagai fallback."


def _open_docx(content: bytes) -> Document:
    return Document(io.BytesIO(content))


def _full_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    parts.append(p.text)
    return "\n".join(parts)


# ── Happy path: ParsedReport ──────────────────────────────────────────────────

def test_stateless_export_with_parsed_returns_200() -> None:
    resp = client.post(
        "/api/v1/export/docx",
        json={"parsed": PARSED_OK.model_dump(), "doc_type": "proposal_skripsi"},
    )
    assert resp.status_code == 200


def test_stateless_export_content_type_is_docx() -> None:
    resp = client.post(
        "/api/v1/export/docx",
        json={"parsed": PARSED_OK.model_dump(), "doc_type": "proposal_skripsi"},
    )
    assert "wordprocessingml" in resp.headers["content-type"]


def test_stateless_export_is_valid_docx() -> None:
    resp = client.post(
        "/api/v1/export/docx",
        json={"parsed": PARSED_OK.model_dump(), "doc_type": "proposal_skripsi"},
    )
    doc = _open_docx(resp.content)
    assert doc is not None


def test_stateless_export_rujukan_count_equals_findings() -> None:
    """F4-4 parity: number of Rujukan table rows == number of findings."""
    resp = client.post(
        "/api/v1/export/docx",
        json={"parsed": PARSED_OK.model_dump(), "doc_type": "proposal_skripsi"},
    )
    doc = _open_docx(resp.content)
    total_data_rows = sum(
        len(table.rows) - 1
        for table in doc.tables
        if len(table.rows) > 1
    )
    assert total_data_rows == len(FINDINGS)


def test_stateless_export_disclaimer_present() -> None:
    resp = client.post(
        "/api/v1/export/docx",
        json={"parsed": PARSED_OK.model_dump(), "doc_type": "proposal_skripsi"},
    )
    doc = _open_docx(resp.content)
    text = _full_text(doc)
    assert text.count("DISCLAIMER") >= 2


# ── raw_output only ───────────────────────────────────────────────────────────

def test_stateless_export_raw_output_only_returns_200() -> None:
    resp = client.post(
        "/api/v1/export/docx",
        json={"raw_output": RAW, "doc_type": "proposal_skripsi"},
    )
    assert resp.status_code == 200


def test_stateless_export_raw_output_in_docx() -> None:
    resp = client.post(
        "/api/v1/export/docx",
        json={"raw_output": RAW, "doc_type": "proposal_skripsi"},
    )
    doc = _open_docx(resp.content)
    text = _full_text(doc)
    assert RAW in text


# ── Empty body → 400 ──────────────────────────────────────────────────────────

def test_stateless_export_empty_body_returns_400() -> None:
    resp = client.post(
        "/api/v1/export/docx",
        json={"doc_type": "proposal_skripsi"},
    )
    assert resp.status_code == 400


def test_stateless_export_empty_body_error_code() -> None:
    resp = client.post(
        "/api/v1/export/docx",
        json={"doc_type": "proposal_skripsi"},
    )
    body = resp.json()
    assert body["error"]["code"] == "EMPTY_BODY"
