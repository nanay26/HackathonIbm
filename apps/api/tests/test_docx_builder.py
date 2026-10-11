"""
Tests for app/services/docx_builder.py (Fase 5 DoD).

Covers:
  - build_docx returns bytes that can be re-opened as a valid .docx
  - Disclaimer at start AND end
  - parse_ok=True: all reviewer sections, findings present, Rujukan verbatim
  - parse_ok=False: raw fallback text included
  - Number of Rujukan cells == number of findings (even null ones)
  - Filename format: Telaah-Review-YYYYMMDD-HHmm.docx
  - No colour (font colour is None/auto)
  - export endpoint returns 200 with correct Content-Type when job is done
"""
from __future__ import annotations

import asyncio
import io
import re

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.models.schemas import (
    DebatePoint,
    DebateType,
    Finding,
    ParsedReport,
    ReviewerKey,
    ReviewerResult,
    ReviewerStatus,
    ReviewJob,
)
from app.services.docx_builder import build_docx, export_filename
from app.services.job_store import _store_holder, get_job_store

# ── Fixtures ──────────────────────────────────────────────────────────────────

FINDINGS_MET = [
    Finding(
        no=1,
        aspect="Desain Penelitian",
        description="Skema validasi tidak dijelaskan.",
        reference='"validasi 10-fold CV" (BAB 3, hlm. 24)',
    ),
    Finding(
        no=2,
        aspect="Reproduksibilitas",
        description="Hyperparameter tidak dicantumkan.",
        reference=None,  # → "Tanpa rujukan"
    ),
]

REVIEWER_OK = ReviewerResult(
    key=ReviewerKey.metodologi,
    status=ReviewerStatus.ok,
    verdict="REVISI MAYOR",
    summary="Metodologi perlu perbaikan.",
    findings=FINDINGS_MET,
    priority_recommendations=["Perbaiki skema validasi.", "Tambahkan hyperparameter."],
)

REVIEWER_MISSING = ReviewerResult(
    key=ReviewerKey.orisinalitas,
    status=ReviewerStatus.missing,
    verdict="",
    summary="",
    findings=[],
    priority_recommendations=[],
)

REVIEWER_FAILED = ReviewerResult(
    key=ReviewerKey.revisi,
    status=ReviewerStatus.failed,
    verdict="",
    summary="",
    findings=[],
    priority_recommendations=[],
)

DEBATE = DebatePoint(
    type=DebateType.penguatan_silang,
    aspect="Validitas Skema Data Split",
    description="Dua reviewer menyoroti isu ini.",
    reviewers=["- Reviewer Metodologi: potensi data leakage."],
    references=[],
)

REPORT_OK = ParsedReport(
    parse_ok=True,
    reviewers=[REVIEWER_OK, REVIEWER_MISSING, REVIEWER_FAILED],
    debate_points=[DEBATE],
)

REPORT_FAILED = ParsedReport(parse_ok=False, reviewers=[], debate_points=[])

RAW_OUTPUT = "Ini adalah output mentah dari Flow 5."


# ── Helper ────────────────────────────────────────────────────────────────────

def _full_text(doc: Document) -> str:
    """Collect all text: paragraphs + table cells."""
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    parts.append(p.text)
    return "\n".join(parts)


def _open_docx(data: bytes) -> Document:
    return Document(io.BytesIO(data))


# ── Bytes output ──────────────────────────────────────────────────────────────

def test_build_docx_returns_bytes() -> None:
    result = build_docx(REPORT_OK)
    assert isinstance(result, bytes)
    assert len(result) > 0


def test_build_docx_is_valid_docx() -> None:
    result = build_docx(REPORT_OK)
    doc = _open_docx(result)  # does not raise
    assert doc is not None


# ── Disclaimer at start AND end ───────────────────────────────────────────────

def test_disclaimer_present_in_document() -> None:
    result = build_docx(REPORT_OK)
    doc = _open_docx(result)
    text = _full_text(doc)
    assert text.count("DISCLAIMER") >= 2, "Disclaimer must appear at least twice (start and end)"


def test_disclaimer_contains_required_text() -> None:
    result = build_docx(REPORT_OK)
    doc = _open_docx(result)
    text = _full_text(doc)
    assert "asistensi akademik" in text.lower()
    assert "dosen pembimbing/penguji resmi" in text.lower()


# ── Structured content ────────────────────────────────────────────────────────

def test_reviewer_section_heading_present() -> None:
    result = build_docx(REPORT_OK)
    doc = _open_docx(result)
    text = _full_text(doc)
    assert "REVIEWER METODOLOGI" in text.upper()


def test_verdict_present() -> None:
    result = build_docx(REPORT_OK)
    doc = _open_docx(result)
    text = _full_text(doc)
    assert "REVISI MAYOR" in text


def test_ringkasan_present() -> None:
    result = build_docx(REPORT_OK)
    doc = _open_docx(result)
    text = _full_text(doc)
    assert "Metodologi perlu perbaikan" in text


def test_aspect_tag_present() -> None:
    result = build_docx(REPORT_OK)
    doc = _open_docx(result)
    text = _full_text(doc)
    assert "ASPEK: Desain Penelitian" in text


def test_description_present() -> None:
    result = build_docx(REPORT_OK)
    doc = _open_docx(result)
    text = _full_text(doc)
    assert "Skema validasi tidak dijelaskan" in text


def test_rujukan_verbatim() -> None:
    """F4-4: Rujukan in docx must be verbatim from parsed result."""
    result = build_docx(REPORT_OK)
    doc = _open_docx(result)
    # Collect all table cell text
    table_text = "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    assert "validasi 10-fold CV" in table_text


def test_tanpa_rujukan_tag_present() -> None:
    """Finding with reference=None → 'Tanpa rujukan' in table cell."""
    result = build_docx(REPORT_OK)
    doc = _open_docx(result)
    table_text = "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    assert "Tanpa rujukan" in table_text


def test_rujukan_count_equals_findings_count() -> None:
    """F4-4: number of Rujukan cells == number of findings."""
    result = build_docx(REPORT_OK)
    doc = _open_docx(result)
    # The findings table has 1 header row + N data rows; each data row col[2] = rujukan
    # Count data rows in all tables
    total_data_rows = sum(
        len(table.rows) - 1  # subtract header row
        for table in doc.tables
        if len(table.rows) > 1
    )
    assert total_data_rows == len(FINDINGS_MET)


def test_priority_recommendations_present() -> None:
    result = build_docx(REPORT_OK)
    doc = _open_docx(result)
    text = _full_text(doc)
    assert "Perbaiki skema validasi" in text


def test_debate_point_present() -> None:
    result = build_docx(REPORT_OK)
    doc = _open_docx(result)
    text = _full_text(doc)
    assert "Validitas Skema Data Split" in text
    assert "PENGUATAN SILANG" in text


def test_failed_missing_reviewer_stub_text() -> None:
    result = build_docx(REPORT_OK)
    doc = _open_docx(result)
    text = _full_text(doc)
    assert "tidak tersedia" in text.lower() or "kesalahan" in text.lower()


# ── parse_ok=False fallback ───────────────────────────────────────────────────

def test_raw_fallback_contains_raw_output() -> None:
    result = build_docx(REPORT_FAILED, raw_output=RAW_OUTPUT)
    doc = _open_docx(result)
    text = _full_text(doc)
    assert RAW_OUTPUT in text


def test_raw_fallback_disclaimer_still_present() -> None:
    result = build_docx(REPORT_FAILED, raw_output=RAW_OUTPUT)
    doc = _open_docx(result)
    text = _full_text(doc)
    assert text.count("DISCLAIMER") >= 2


def test_raw_fallback_no_reviewer_tables() -> None:
    """When parse_ok=False, no structured tables should be present."""
    result = build_docx(REPORT_FAILED, raw_output=RAW_OUTPUT)
    doc = _open_docx(result)
    assert len(doc.tables) == 0


# ── Filename format ───────────────────────────────────────────────────────────

def test_export_filename_format() -> None:
    name = export_filename()
    assert name.startswith("Telaah-Review-")
    assert name.endswith(".docx")
    # Date part: YYYYMMDD-HHmm
    pattern = r"^Telaah-Review-\d{8}-\d{4}\.docx$"
    assert re.match(pattern, name), f"Unexpected filename: {name}"


# ── No colour ─────────────────────────────────────────────────────────────────

def test_no_font_colour_set() -> None:
    """F4-2: no colour — all runs should have no explicit font colour."""
    result = build_docx(REPORT_OK)
    doc = _open_docx(result)
    for para in doc.paragraphs:
        for run in para.runs:
            # RGBColor None means no explicit colour (auto/default)
            assert run.font.color.rgb is None or run.font.color.type is None, (
                f"Run '{run.text[:30]}' has explicit color set"
            )


# ── Export endpoint integration ───────────────────────────────────────────────

@pytest.fixture(autouse=True)
def reset_job_store() -> None:
    _store_holder.clear()


client = TestClient(app, raise_server_exceptions=False)
VALID_TEXT = "a" * 2000


def test_export_returns_200_with_docx_content_type() -> None:
    """Inject a done job into the store, then GET export.docx → 200."""
    from app.models.schemas import JobStatus

    store = get_job_store()
    job = ReviewJob(
        status=JobStatus.done,
        raw_output="output",
        parsed=ParsedReport(parse_ok=False),
    )
    loop = asyncio.new_event_loop()
    try:
        loop.run_until_complete(store.create(job))
    finally:
        loop.close()

    resp = client.get(f"/api/v1/reviews/{job.id}/export.docx")
    assert resp.status_code == 200
    assert "wordprocessingml" in resp.headers["content-type"]


def test_export_returns_410_for_missing_job() -> None:
    resp = client.get("/api/v1/reviews/no-such-id/export.docx")
    assert resp.status_code == 410


def test_export_returns_409_for_queued_job() -> None:
    create = client.post(
        "/api/v1/reviews",
        json={"doc_type": "proposal_skripsi", "text": VALID_TEXT},
    )
    review_id = create.json()["review_id"]
    resp = client.get(f"/api/v1/reviews/{review_id}/export.docx")
    assert resp.status_code in {409, 200}  # 200 if background task already finished
